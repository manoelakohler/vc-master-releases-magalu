"""Geração da planilha — o artefato final principal.

A planilha precisa sustentar a pergunta de quem não acompanhou a execução: "de
onde veio este número?". A resposta está nela mesma: de qualquer valor exibido
dá para chegar a uma linha de Evidências com documento, página e trecho.

Duas decisões que parecem detalhes e não são:

- **Nenhuma fórmula viva.** É registro auditável, não modelo recalculável: uma
  fórmula que recalcula pode divergir do que foi extraído e auditado.
- **Texto literal gravado como texto.** O Excel converte `9.856,4` em data, em
  número anglófono ou em coisa pior conforme o locale de quem abrir o arquivo —
  e a coluna que existe para provar o documento é justamente a que se perde.
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from magalu_releases.auditoria.checks import verificacao
from magalu_releases.models import VerificacaoAuditoria
from magalu_releases.periodos import PeriodoIndeterminado, interpretar_periodo
from magalu_releases.saida.resumo import contar_periodos_analisados
from magalu_releases.vocabularios import Severidade

ABAS = ("Resumo", "Comparativo", "Evidências", "Documentos", "Pendências", "Auditoria")

# Os quatro checks de planilha da skill, na ordem em que aparecem na aba.
_CHECKS_PLANILHA = (
    ("xls_abas", "As seis abas existem, com nomes e ordem corretos"),
    ("xls_colunas_periodo", "Comparativo tem N colunas de período em ordem crescente"),
    ("xls_texto_preservado", "valor_original e trecho_fonte gravados como texto"),
    ("xls_reabertura", "Arquivo reabre e as verificações passam sobre o conteúdo lido"),
)

_FONTE = "Arial"
_FONTE_CABECALHO = Font(name=_FONTE, bold=True, color="FFFFFF")
_FONTE_CORPO = Font(name=_FONTE)
_FONTE_TITULO = Font(name=_FONTE, bold=True, size=13)
_FUNDO_CABECALHO = PatternFill("solid", fgColor="1F3864")
_TEXTO = "@"

_COLUNAS_EVIDENCIAS = (
    "fato_id", "metrica_id", "metrica_rotulo", "segmento", "base", "periodicidade",
    "tipo_valor", "periodo_canonico", "periodo_rotulo", "valor_original",
    "valor_normalizado", "unidade", "documento_id", "pagina", "trecho_fonte",
    "confianca", "flags", "revisao_humana", "motivo",
)
_COLUNAS_DOCUMENTOS = (
    "documento_id", "titulo", "tipo", "periodo_canonico", "periodo_rotulo", "url_origem",
    "nome_servidor", "data_publicacao", "arquivo_local", "bytes", "sha256", "paginas",
    "textual", "baixado_em", "motivo_descarte",
)
_COLUNAS_PENDENCIAS = (
    "pendencia_id", "tipo", "severidade", "descricao", "referencias",
    "valores_conflitantes", "acao_sugerida", "status",
)
_COLUNAS_AUDITORIA = (
    "check_id", "descricao", "esperado", "obtido", "resultado", "severidade", "detalhe",
)

# Colunas cujo conteúdo é prova documental e não pode ser reinterpretado.
_COLUNAS_LITERAIS = frozenset(
    {"valor_original", "trecho_fonte", "periodo_rotulo", "periodo_canonico", "sha256",
     "nome_servidor",
     "fato_id", "documento_id", "serie_id", "pendencia_id", "check_id"}
)


def _escrever_cabecalho(aba, colunas, linha=1):
    for indice, nome in enumerate(colunas, start=1):
        celula = aba.cell(row=linha, column=indice, value=nome)
        celula.font = _FONTE_CABECALHO
        celula.fill = _FUNDO_CABECALHO
        celula.alignment = Alignment(vertical="center", wrap_text=False)
    aba.freeze_panes = aba.cell(row=linha + 1, column=1)


def _escrever_celula(aba, linha, coluna, valor, nome_coluna=""):
    """Grava preservando o tipo. `None` fica vazio — nunca vira zero."""
    celula = aba.cell(row=linha, column=coluna)
    celula.font = _FONTE_CORPO
    if valor is None:
        celula.value = None
        return celula
    if nome_coluna in _COLUNAS_LITERAIS or isinstance(valor, str):
        celula.value = str(valor)
        celula.number_format = _TEXTO
    else:
        celula.value = valor
    return celula


def _ajustar_larguras(aba, colunas, maximo=60):
    for indice, nome in enumerate(colunas, start=1):
        largura = min(max(len(str(nome)) + 4, 12), maximo)
        aba.column_dimensions[get_column_letter(indice)].width = largura


def _valor_enum(valor):
    return valor.value if hasattr(valor, "value") else valor


def gerar_excel(
    *,
    caminho: Path | str,
    n_pedido: int,
    periodos,
    documentos,
    fatos,
    series,
    variacoes,
    pendencias,
    resumo_texto: str,
    auditoria,
    fonte: str,
    run_id: str,
) -> Path:
    """Escreve a planilha completa e devolve o caminho gravado."""
    periodos = tuple(periodos)
    destino = Path(caminho)
    destino.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)
    for nome in ABAS:
        wb.create_sheet(nome)

    rotulos = _rotulos_por_periodo(documentos, periodos)

    _aba_resumo(wb["Resumo"], n_pedido, periodos, rotulos, documentos, pendencias,
                resumo_texto, auditoria, fonte, run_id, series, variacoes)
    _aba_comparativo(wb["Comparativo"], series, variacoes, periodos, rotulos)
    _aba_evidencias(wb["Evidências"], fatos)
    _aba_documentos(wb["Documentos"], documentos)
    _aba_pendencias(wb["Pendências"], pendencias)
    _aba_auditoria(wb["Auditoria"], auditoria)

    wb.save(destino)
    return destino


def _rotulos_por_periodo(documentos, periodos) -> dict[str, str]:
    """Rótulo como o documento escreveu (2T25), com o canônico como reserva."""
    mapa = {}
    for documento in documentos:
        if documento.periodo:
            mapa.setdefault(documento.periodo.canonico, documento.periodo.rotulo)
    return {c: mapa.get(c, c) for c in periodos}


def _aba_resumo(aba, n_pedido, periodos, rotulos, documentos, pendencias, resumo_texto,
                auditoria, fonte, run_id, series, variacoes):
    aba.cell(row=1, column=1, value="Análise de releases de resultados").font = _FONTE_TITULO

    n_obtido = contar_periodos_analisados(periodos, documentos)
    falhas = [v for v in auditoria if _valor_enum(v.resultado) == "FAIL"]

    cabecalho = [
        ("Empresa", "Magazine Luiza"),
        ("Fonte", fonte),
        ("Execução", run_id),
        ("Gerado em", datetime.now().isoformat(timespec="seconds")),
        ("N pedido", n_pedido),
        ("N obtido", n_obtido),
        ("Períodos analisados", ", ".join(rotulos[c] for c in periodos)),
        ("Documentos utilizados", len(documentos)),
        ("Séries construídas", len(series)),
        ("Pendências abertas", len(pendencias)),
        ("Verificações de auditoria", len(auditoria)),
        ("Falhas de auditoria", len(falhas)),
    ]
    linha = 3
    for rotulo, valor in cabecalho:
        aba.cell(row=linha, column=1, value=rotulo).font = Font(name=_FONTE, bold=True)
        _escrever_celula(aba, linha, 2, valor)
        linha += 1

    if n_obtido < n_pedido:
        aviso = aba.cell(
            row=linha,
            column=1,
            value=f"ATENÇÃO: foram pedidos {n_pedido} releases e obtidos {n_obtido}.",
        )
        aviso.font = Font(name=_FONTE, bold=True, color="C00000")
        linha += 1

    linha += 1
    aba.cell(row=linha, column=1, value="Resumo executivo").font = _FONTE_TITULO
    linha += 1
    for paragrafo in (resumo_texto or "").splitlines() or [""]:
        celula = _escrever_celula(aba, linha, 1, paragrafo)
        celula.alignment = Alignment(wrap_text=True, vertical="top")
        linha += 1

    aba.column_dimensions["A"].width = 34
    aba.column_dimensions["B"].width = 70


def _aba_comparativo(aba, series, variacoes, periodos, rotulos):
    colunas = [
        "serie_id", "metrica_id", "metrica_rotulo", "segmento", "base", "periodicidade",
        "tipo_valor", "unidade_serie",
    ]
    colunas += [rotulos[c] for c in periodos]
    colunas += [
        "variacao_abs", "variacao_pct", "variacao_pp", "base_comparacao", "calculada",
        "motivo_nao_calculada", "confianca_minima", "evidencias", "revisao_humana",
    ]
    _escrever_cabecalho(aba, colunas)
    _ajustar_larguras(aba, colunas)

    ultima_variacao = {}
    for variacao in variacoes:
        ultima_variacao[variacao.serie_id] = variacao

    for indice, serie in enumerate(series, start=2):
        valores = [
            serie.serie_id, serie.metrica_id, serie.metrica_rotulo,
            _valor_enum(serie.segmento), _valor_enum(serie.base),
            _valor_enum(serie.periodicidade), _valor_enum(serie.tipo_valor),
            serie.unidade_serie,
        ]
        por_periodo = {p.periodo_canonico: p for p in serie.pontos}
        for canonico in periodos:
            ponto = por_periodo.get(canonico)
            valores.append(ponto.valor_normalizado if ponto else None)

        v = ultima_variacao.get(serie.serie_id)
        valores += [
            v.variacao_abs if v else None,
            v.variacao_pct if v else None,
            v.variacao_pp if v else None,
            _valor_enum(v.base_comparacao) if v and v.base_comparacao else None,
            v.calculada if v else None,
            v.motivo_nao_calculada if v else None,
            _valor_enum(serie.confianca_minima),
            sum(1 for p in serie.pontos if p.fato_id),
            serie.revisao_humana,
        ]
        for coluna, (nome, valor) in enumerate(zip(colunas, valores), start=1):
            _escrever_celula(aba, indice, coluna, valor, nome)
    aba.auto_filter.ref = aba.dimensions


def _aba_evidencias(aba, fatos):
    _escrever_cabecalho(aba, _COLUNAS_EVIDENCIAS)
    _ajustar_larguras(aba, _COLUNAS_EVIDENCIAS)
    for indice, fato in enumerate(fatos, start=2):
        valores = {
            "fato_id": fato.fato_id,
            "metrica_id": fato.metrica_id,
            "metrica_rotulo": fato.metrica_rotulo,
            "segmento": _valor_enum(fato.segmento),
            "base": _valor_enum(fato.base),
            "periodicidade": _valor_enum(fato.periodicidade),
            "tipo_valor": _valor_enum(fato.tipo_valor),
            "periodo_canonico": fato.periodo.canonico,
            "periodo_rotulo": fato.periodo.rotulo,
            "valor_original": fato.valor_original,
            "valor_normalizado": fato.valor_normalizado,
            "unidade": fato.unidade,
            "documento_id": fato.documento_id,
            "pagina": fato.pagina,
            "trecho_fonte": fato.trecho_fonte,
            "confianca": _valor_enum(fato.confianca),
            "flags": ", ".join(_valor_enum(g) for g in fato.flags) or None,
            "revisao_humana": fato.revisao_humana,
            "motivo": fato.motivo,
        }
        for coluna, nome in enumerate(_COLUNAS_EVIDENCIAS, start=1):
            _escrever_celula(aba, indice, coluna, valores[nome], nome)
    aba.auto_filter.ref = aba.dimensions


def _aba_documentos(aba, documentos):
    _escrever_cabecalho(aba, _COLUNAS_DOCUMENTOS)
    _ajustar_larguras(aba, _COLUNAS_DOCUMENTOS)
    for indice, d in enumerate(documentos, start=2):
        valores = {
            "documento_id": d.documento_id,
            "titulo": d.titulo,
            "tipo": _valor_enum(d.tipo),
            "periodo_canonico": d.periodo.canonico if d.periodo else None,
            "periodo_rotulo": d.periodo.rotulo if d.periodo else None,
            "url_origem": d.url_origem,
            "nome_servidor": d.nome_servidor,
            "data_publicacao": d.data_publicacao,
            "arquivo_local": d.arquivo_local,
            "bytes": d.bytes,
            "sha256": d.sha256,
            "paginas": d.paginas,
            "textual": d.textual,
            "baixado_em": d.baixado_em,
            "motivo_descarte": d.motivo_descarte,
        }
        for coluna, nome in enumerate(_COLUNAS_DOCUMENTOS, start=1):
            _escrever_celula(aba, indice, coluna, valores[nome], nome)


def _aba_pendencias(aba, pendencias):
    _escrever_cabecalho(aba, _COLUNAS_PENDENCIAS)
    _ajustar_larguras(aba, _COLUNAS_PENDENCIAS)

    if not pendencias:
        # Aba vazia é ambígua: não se sabe se nada foi encontrado ou se a
        # verificação não rodou. A declaração explícita remove a dúvida.
        _escrever_celula(aba, 2, 1, "—", "pendencia_id")
        _escrever_celula(aba, 2, 4, "Nenhuma pendência registrada nesta execução.")
        return

    ordem = {"alta": 0, "media": 1, "baixa": 2}
    ordenadas = sorted(pendencias, key=lambda p: ordem.get(_valor_enum(p.severidade), 9))
    for indice, p in enumerate(ordenadas, start=2):
        valores = {
            "pendencia_id": p.pendencia_id,
            "tipo": _valor_enum(p.tipo),
            "severidade": _valor_enum(p.severidade),
            "descricao": p.descricao,
            "referencias": ", ".join(p.referencias) or None,
            "valores_conflitantes": " | ".join(p.valores_conflitantes) or None,
            "acao_sugerida": p.acao_sugerida or None,
            "status": p.status,
        }
        for coluna, nome in enumerate(_COLUNAS_PENDENCIAS, start=1):
            _escrever_celula(aba, indice, coluna, valores[nome], nome)
    aba.auto_filter.ref = aba.dimensions


def _aba_auditoria(aba, verificacoes):
    _escrever_cabecalho(aba, _COLUNAS_AUDITORIA)
    _ajustar_larguras(aba, _COLUNAS_AUDITORIA)
    vermelho = Font(name=_FONTE, bold=True, color="C00000")
    for indice, v in enumerate(verificacoes, start=2):
        valores = {
            "check_id": v.check_id,
            "descricao": v.descricao,
            "esperado": v.esperado,
            "obtido": v.obtido,
            "resultado": _valor_enum(v.resultado),
            "severidade": _valor_enum(v.severidade),
            "detalhe": v.detalhe or None,
        }
        for coluna, nome in enumerate(_COLUNAS_AUDITORIA, start=1):
            celula = _escrever_celula(aba, indice, coluna, valores[nome], nome)
            if nome == "resultado" and valores[nome] == "FAIL":
                celula.font = vermelho


def validar_planilha(
    caminho: Path | str, *, n_periodos: int, periodos=()
) -> tuple[VerificacaoAuditoria, ...]:
    """Reabre o arquivo e confere o conteúdo, como verificações auditáveis.

    Gravar sem erro não prova que a planilha está certa — só a releitura prova.
    E o resultado dessa releitura é auditoria: vira linha na aba Auditoria, como
    qualquer outro check, inclusive quando passa.
    """
    destino = Path(caminho)
    if not destino.is_file():
        return _todas_falham(f"arquivo não encontrado: {destino}")

    try:
        wb = load_workbook(destino)
    except Exception as erro:  # arquivo corrompido é falha de reabertura
        return _todas_falham(f"arquivo não pôde ser reaberto: {erro}")

    abas_ok = wb.sheetnames == list(ABAS)
    checks = [
        verificacao(
            "xls_abas",
            "As seis abas existem, com nomes e ordem corretos",
            list(ABAS),
            wb.sheetnames,
            abas_ok,
            Severidade.ALTA,
        )
    ]
    if not abas_ok:
        # Sem as abas certas, os demais checks não têm onde olhar.
        checks.extend(
            _falha(check_id, descricao, "estrutura de abas inválida")
            for check_id, descricao in _CHECKS_PLANILHA[1:]
        )
        return tuple(checks)

    comparativo = wb["Comparativo"]
    cabecalhos = [c.value for c in comparativo[1]]
    colunas_periodo = [str(h) for h in cabecalhos if h and _parece_periodo(str(h))]
    detalhes_periodo = []
    if len(colunas_periodo) != n_periodos:
        detalhes_periodo.append(
            f"{len(colunas_periodo)} coluna(s) de período para {n_periodos} período(s)"
        )
    if not _em_ordem_crescente(colunas_periodo):
        detalhes_periodo.append(f"fora de ordem crescente: {colunas_periodo}")
    checks.append(
        verificacao(
            "xls_colunas_periodo",
            "Comparativo tem N colunas de período em ordem crescente",
            f"{n_periodos} colunas em ordem crescente",
            f"{len(colunas_periodo)} coluna(s)",
            not detalhes_periodo,
            Severidade.ALTA,
            "; ".join(detalhes_periodo),
        )
    )

    evidencias = wb["Evidências"]
    colunas_evidencias = [c.value for c in evidencias[1]]
    detalhes_texto = [
        f"Evidências sem a coluna {obrigatoria!r}"
        for obrigatoria in ("documento_id", "pagina", "trecho_fonte", "valor_original")
        if obrigatoria not in colunas_evidencias
    ]
    # Toda linha, não só a primeira: o trecho é a prova documental, e uma célula
    # reinterpretada pelo Excel deixa de provar o que estava no release.
    for nome in ("valor_original", "trecho_fonte"):
        if nome not in colunas_evidencias:
            continue
        coluna = colunas_evidencias.index(nome) + 1
        reformatadas = [
            linha
            for linha in range(2, evidencias.max_row + 1)
            if evidencias.cell(row=linha, column=coluna).value is not None
            and evidencias.cell(row=linha, column=coluna).number_format != _TEXTO
        ]
        if reformatadas:
            detalhes_texto.append(
                f"{nome} fora do formato texto na(s) linha(s) "
                + ", ".join(str(linha) for linha in reformatadas[:10])
            )
    checks.append(
        verificacao(
            "xls_texto_preservado",
            "valor_original e trecho_fonte gravados como texto",
            "0 células reinterpretadas",
            f"{len(detalhes_texto)} problema(s)",
            not detalhes_texto,
            Severidade.ALTA,
            "; ".join(detalhes_texto),
        )
    )

    detalhes_reabertura = []
    if "documento_id" in colunas_evidencias:
        ids_documentos = _coluna_como_conjunto(wb["Documentos"], "documento_id")
        coluna = colunas_evidencias.index("documento_id") + 1
        for linha in range(2, evidencias.max_row + 1):
            valor = evidencias.cell(row=linha, column=coluna).value
            if valor and valor not in ids_documentos:
                detalhes_reabertura.append(
                    f"Evidências referenciam documento inexistente: {valor!r}"
                )
    if wb["Pendências"].max_row < 2:
        detalhes_reabertura.append(
            "aba Pendências vazia: declare explicitamente quando não houver"
        )
    if wb["Auditoria"].max_row < 2:
        detalhes_reabertura.append(
            "aba Auditoria sem verificações: a auditoria não foi registrada"
        )
    checks.append(
        verificacao(
            "xls_reabertura",
            "Arquivo reabre e as verificações passam sobre o conteúdo lido",
            "0 problemas na releitura",
            f"{len(detalhes_reabertura)} problema(s)",
            not detalhes_reabertura,
            Severidade.ALTA,
            "; ".join(dict.fromkeys(detalhes_reabertura)),
        )
    )

    return tuple(checks)


def _falha(check_id: str, descricao: str, motivo: str) -> VerificacaoAuditoria:
    return verificacao(check_id, descricao, "verificação executada", motivo, False, Severidade.ALTA)


def _todas_falham(motivo: str) -> tuple[VerificacaoAuditoria, ...]:
    """Planilha ilegível reprova as quatro: nenhuma delas chegou a rodar."""
    return tuple(
        _falha(check_id, descricao, motivo) for check_id, descricao in _CHECKS_PLANILHA
    )


_CANONICO = re.compile(r"^(\d{4})-(?:Q([1-4])|(\d{1,2})M|FY)$")


def _chave_do_cabecalho(texto: str) -> tuple[int, int] | None:
    """Chave de ordenação do cabeçalho, venha ele como rótulo ou como canônico.

    A coluna leva o rótulo do documento (`2T26`) quando algum documento o
    trouxe, e o canônico (`2026-Q2`) como reserva. Entender só uma das duas
    formas faz o check passar sem verificar nada na outra.
    """
    try:
        return interpretar_periodo(texto).chave_ordenacao
    except PeriodoIndeterminado:
        pass

    achado = _CANONICO.match(texto.strip())
    if not achado:
        return None
    ano = int(achado.group(1))
    if achado.group(2):
        return (ano, int(achado.group(2)))
    if achado.group(3):
        return (ano, int(achado.group(3)) // 3)
    return (ano, 4)


def _em_ordem_crescente(rotulos) -> bool:
    """Ordem pelo período fiscal, não pela string: 4T25 vem antes de 1T26."""
    chaves = []
    for rotulo in rotulos:
        chave = _chave_do_cabecalho(rotulo)
        if chave is None:
            return True  # cabeçalho que não é período não diz nada sobre a ordem
        chaves.append(chave)
    return chaves == sorted(chaves)


def _parece_periodo(texto: str) -> bool:
    import re

    return bool(re.fullmatch(r"[1-4]T\d{2,4}|\d{4}-(?:Q[1-4]|\d{1,2}M|FY)", texto.strip()))



def _coluna_como_conjunto(aba, nome_coluna) -> set:
    cabecalhos = [c.value for c in aba[1]]
    if nome_coluna not in cabecalhos:
        return set()
    coluna = cabecalhos.index(nome_coluna) + 1
    return {
        aba.cell(row=linha, column=coluna).value
        for linha in range(2, aba.max_row + 1)
        if aba.cell(row=linha, column=coluna).value
    }
